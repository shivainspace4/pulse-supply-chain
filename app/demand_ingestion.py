
import csv
import io
from datetime import date
from typing import Annotated, Any

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel, Field, ValidationError
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from app.database import SessionLocal
from app.db_models import DemandRow, LocationRow, ProductRow


router = APIRouter(prefix="/demand", tags=["Demand Ingestion"])


def get_db():
    with SessionLocal() as session:
        yield session


class DemandRecordRequest(BaseModel):
    sku: str = Field(min_length=1, max_length=50)
    depot_id: str = Field(min_length=1, max_length=50)
    demand_date: date
    forecast_quantity: int = Field(ge=0)
    actual_quantity: int | None = Field(default=None, ge=0)


def validate_references(record: DemandRecordRequest, db: Session):
    product = db.get(ProductRow, record.sku)
    if product is None:
        raise HTTPException(
            status_code=422,
            detail=f"Unknown SKU: {record.sku}",
        )

    location = db.get(LocationRow, record.depot_id)
    if location is None or location.location_type != "DEPOT":
        raise HTTPException(
            status_code=422,
            detail=f"Unknown depot: {record.depot_id}",
        )


def existing_record(
    db: Session,
    sku: str,
    depot_id: str,
    demand_date: date,
) -> bool:
    statement = select(DemandRow.demand_id).where(
        DemandRow.sku == sku,
        DemandRow.depot_id == depot_id,
        DemandRow.demand_date == demand_date,
    )
    return db.scalar(statement) is not None


def serialize_record(record: DemandRow) -> dict[str, Any]:
    return {
        "demand_id": record.demand_id,
        "sku": record.sku,
        "depot_id": record.depot_id,
        "demand_date": record.demand_date.isoformat(),
        "forecast_quantity": record.forecast_quantity,
        "actual_quantity": record.actual_quantity,
    }


@router.post("/records", status_code=201)
def create_demand_record(
    request: DemandRecordRequest,
    db: Session = Depends(get_db),
):
    validate_references(request, db)

    if existing_record(
        db, request.sku, request.depot_id, request.demand_date
    ):
        raise HTTPException(
            status_code=409,
            detail=(
                "A demand record already exists for this SKU, "
                "depot and date."
            ),
        )

    record = DemandRow(
        sku=request.sku,
        depot_id=request.depot_id,
        demand_date=request.demand_date,
        forecast_quantity=request.forecast_quantity,
        actual_quantity=request.actual_quantity,
    )

    try:
        db.add(record)
        db.commit()
        db.refresh(record)
    except IntegrityError:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="The record conflicts with existing database data.",
        )

    return serialize_record(record)


def parse_upload(filename: str, content: bytes) -> list[dict[str, Any]]:
    name = filename.lower()

    try:
        if name.endswith(".csv"):
            text = content.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            if not reader.fieldnames:
                raise ValueError("CSV is missing its header row.")
            return [
                {str(k).strip().lower(): v for k, v in row.items() if k}
                for row in reader
            ]

        if name.endswith(".xlsx"):
            workbook = load_workbook(
                io.BytesIO(content),
                read_only=True,
                data_only=True,
            )
            try:
                sheet = workbook.active
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, None)

                if not headers:
                    raise ValueError("Excel sheet is empty.")

                normalized = [
                    str(value).strip().lower() if value is not None else ""
                    for value in headers
                ]

                result = []
                for values in rows:
                    if all(value is None for value in values):
                        continue
                    result.append(dict(zip(normalized, values)))

                return result
            finally:
                workbook.close()

        raise ValueError("Only .csv and .xlsx files are supported.")

    except (UnicodeDecodeError, csv.Error) as exc:
        raise ValueError(f"Could not parse the file: {exc}") from exc


def normalize_row(
    row: dict[str, Any],
    row_number: int,
) -> DemandRecordRequest:
    required = {
        "sku",
        "depot_id",
        "demand_date",
        "forecast_quantity",
    }
    missing = required - set(row)

    if missing:
        raise ValueError(
            f"Missing required columns: {', '.join(sorted(missing))}"
        )

    values = dict(row)

    # Treat blank actuals as missing, not as zero.
    if values.get("actual_quantity") in ("", None):
        values["actual_quantity"] = None

    for field in ("sku", "depot_id"):
        if values.get(field) is not None:
            values[field] = str(values[field]).strip()

    # CSV quantities arrive as strings; Excel quantities may be numeric.
    for field in ("forecast_quantity", "actual_quantity"):
        value = values.get(field)
        if value is None:
            continue

        if isinstance(value, bool):
            raise ValueError(f"{field} must be a whole number.")

        if isinstance(value, float) and not value.is_integer():
            raise ValueError(f"{field} must be a whole number.")

        try:
            values[field] = int(value)
        except (ValueError, TypeError, OverflowError) as exc:
            raise ValueError(
                f"{field} must be a whole number."
            ) from exc

    try:
        return DemandRecordRequest.model_validate(values)
    except ValidationError as exc:
        messages = [
            f"{'.'.join(str(part) for part in error['loc'])}: "
            f"{error['msg']}"
            for error in exc.errors()
        ]
        raise ValueError("; ".join(messages)) from exc


@router.post("/import")
async def import_demand_file(
    file: Annotated[UploadFile, File()],
    db: Session = Depends(get_db),
):
    filename = file.filename or ""

    if not filename.lower().endswith((".csv", ".xlsx")):
        raise HTTPException(
            status_code=400,
            detail="Upload a .csv or .xlsx file.",
        )

    content = await file.read()

    if not content:
        raise HTTPException(status_code=400, detail="The file is empty.")

    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(
            status_code=413,
            detail="Maximum upload size is 10 MB.",
        )

    try:
        rows = parse_upload(filename, content)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    if not rows:
        raise HTTPException(
            status_code=400,
            detail="The file contains no data rows.",
        )

    if len(rows) > 10000:
        raise HTTPException(
            status_code=400,
            detail="Maximum import size is 10,000 rows.",
        )

    valid_records = []
    errors = []
    seen_keys = set()

    for index, row in enumerate(rows, start=2):
        try:
            record = normalize_row(row, index)
            validate_references(record, db)

            key = (record.sku, record.depot_id, record.demand_date)

            if key in seen_keys:
                raise ValueError(
                    "Duplicate SKU, depot and date within this file."
                )

            if existing_record(
                db, record.sku, record.depot_id, record.demand_date
            ):
                raise ValueError(
                    "A record already exists for this SKU, depot and date."
                )

            seen_keys.add(key)
            valid_records.append(record)

        except (ValueError, HTTPException) as exc:
            message = (
                exc.detail if isinstance(exc, HTTPException) else str(exc)
            )
            errors.append({"row": index, "error": message})

    # Atomic import: if any row fails validation, insert nothing.
    if errors:
        return {
            "status": "validation_failed",
            "rows_received": len(rows),
            "rows_imported": 0,
            "rows_rejected": len(errors),
            "errors": errors,
            "message": "Fix the rejected rows and upload the file again.",
        }

    try:
        db.add_all([
            DemandRow(
                sku=record.sku,
                depot_id=record.depot_id,
                demand_date=record.demand_date,
                forecast_quantity=record.forecast_quantity,
                actual_quantity=record.actual_quantity,
            )
            for record in valid_records
        ])
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail=(
                "Import conflicted with database constraints. "
                "No records were imported; check for concurrent imports."
            ),
        ) from exc

    return {
        "status": "success",
        "rows_received": len(rows),
        "rows_imported": len(valid_records),
        "rows_rejected": 0,
        "errors": [],
        "message": "All demand records were imported successfully.",
    }
